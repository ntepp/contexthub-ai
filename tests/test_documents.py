import io
import uuid
from pathlib import Path

import docx

from src.config import get_settings
from src.db.models import Document
from src.ingestion.pipeline import ingest_document
from tests.conftest import upload
from tests.fakes import make_pdf

PDF = make_pdf(
    [
        ["TITRE VII", "Dispositions finales", "Article 93", "Transposition",
         "Les Etats membres appliquent ces dispositions a partir du 3 janvier 2017."],
        ["Article 94", "Entree en vigueur",
         "La presente directive entre en vigueur le vingtieme jour suivant sa publication."],
    ]
)


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Politique de meilleure execution", level=1)
    document.add_paragraph("La societe de gestion selectionne les intermediaires selon des criteres documentes.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_upload_is_indexed_with_pages_and_sections(admin, collection, store):
    resp = upload(admin, collection["id"], "directive.pdf", PDF)
    assert resp.status_code == 202, resp.text
    doc_id = resp.json()["id"]

    doc = admin.get(f"/documents/{doc_id}").json()
    assert doc["status"] == "ready", doc
    assert doc["page_count"] == 2
    assert doc["chunk_count"] == store.count_document(uuid.UUID(doc_id)) > 0
    assert doc["progress"] == {"current": doc["chunk_count"], "total": doc["chunk_count"]}

    chunks = sorted(store.chunks.values(), key=lambda c: c.chunk_index)
    first = chunks[0]
    assert first.metadata["page"] == 1  # 1-based, as displayed to users
    assert first.metadata["section"] == "Article 93 — Transposition"
    assert first.metadata["section_path"][0] == "Titre VII — Dispositions finales"
    assert first.content.startswith("[Titre VII — Dispositions finales › Article 93 — Transposition]")
    assert any(c.metadata["page"] == 2 and c.metadata["section"].startswith("Article 94") for c in chunks)


def test_file_is_stored_under_generated_name_and_viewable(admin, collection):
    doc = upload(admin, collection["id"], "../../evil name.pdf", PDF).json()
    assert doc["filename"] == "evil name.pdf"
    stored = list(Path(get_settings().upload_dir).rglob(f"{doc['id']}.pdf"))
    assert len(stored) == 1 and collection["id"] in stored[0].as_posix()

    view = admin.get(f"/documents/{doc['id']}/file")
    assert view.status_code == 200
    assert view.content == PDF
    assert view.headers["content-type"] == "application/pdf"
    assert "inline" in view.headers["content-disposition"]


def test_duplicate_upload_is_rejected(admin, collection):
    first = upload(admin, collection["id"], "a.pdf", PDF)
    second = upload(admin, collection["id"], "copie.pdf", PDF)
    assert second.status_code == 409
    assert second.json()["detail"]["document_id"] == first.json()["id"]


def test_same_file_allowed_in_another_collection(admin, collection):
    other = admin.post("/collections", json={"name": "Autre"}).json()
    assert upload(admin, collection["id"], "a.pdf", PDF).status_code == 202
    assert upload(admin, other["id"], "a.pdf", PDF).status_code == 202


def test_upload_validation(admin, collection, monkeypatch):
    assert upload(admin, collection["id"], "x.exe", b"MZ...", "application/octet-stream").status_code == 415
    assert upload(admin, collection["id"], "fake.pdf", b"PK\x03\x04 not a pdf").status_code == 400
    assert upload(admin, collection["id"], "empty.txt", b"", "text/plain").status_code == 400
    assert upload(admin, collection["id"], "bin.txt", b"abc\x00def", "text/plain").status_code == 400
    monkeypatch.setattr(get_settings(), "upload_max_mb", 0)
    assert upload(admin, collection["id"], "big.pdf", PDF).status_code == 413


def test_docx_txt_md_html_are_supported(admin, collection, store):
    files = [
        ("politique.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("note.txt", "Note interne sur la gestion des conflits d interets du groupe.".encode(), "text/plain"),
        ("guide.md", "# Guide\n\nProcedure d escalade des alertes de conformite.".encode(), "text/markdown"),
        ("page.html", b"<html><body><h2>FAQ</h2><p>Les reclamations clients sont traitees sous dix jours.</p>"
                      b"<script>alert(1)</script></body></html>", "text/html"),
    ]
    for name, content, mime in files:
        resp = upload(admin, collection["id"], name, content, mime)
        assert resp.status_code == 202, (name, resp.text)
        assert admin.get(f"/documents/{resp.json()['id']}").json()["status"] == "ready", name
    sections = {c.metadata["section"] for c in store.chunks.values()}
    assert {"Politique de meilleure execution", "Guide", "FAQ"} <= sections
    assert not any("alert(1)" in c.content for c in store.chunks.values())


def test_unreadable_pdf_fails_with_explanation(admin, collection, store):
    resp = upload(admin, collection["id"], "broken.pdf", b"%PDF-1.4\ngarbage that is not a pdf")
    doc = admin.get(f"/documents/{resp.json()['id']}").json()
    assert doc["status"] == "failed"
    assert doc["error_message"]
    assert store.count_document(uuid.UUID(doc["id"])) == 0


def test_reingestion_is_idempotent(admin, collection, store):
    doc_id = uuid.UUID(upload(admin, collection["id"], "a.pdf", PDF).json()["id"])
    before = store.count_document(doc_id)
    ingest_document(doc_id)
    ingest_document(doc_id)
    assert store.count_document(doc_id) == before


def test_delete_document_removes_chunks_file_and_row(admin, collection, store, db):
    doc = upload(admin, collection["id"], "a.pdf", PDF).json()
    path = next(Path(get_settings().upload_dir).rglob(f"{doc['id']}.pdf"))
    assert admin.delete(f"/documents/{doc['id']}").status_code == 200
    assert store.count_document(uuid.UUID(doc["id"])) == 0
    assert not path.exists()
    assert db.get(Document, uuid.UUID(doc["id"])) is None
    actions = [e["action"] for e in admin.get("/admin/audit").json()["items"]]
    assert "document.delete" in actions


def test_delete_collection_cascades(admin, collection, store, db):
    doc = upload(admin, collection["id"], "a.pdf", PDF).json()
    assert admin.delete(f"/collections/{collection['id']}").status_code == 200
    assert not store.chunks
    assert db.get(Document, uuid.UUID(doc["id"])) is None
    assert not (Path(get_settings().upload_dir) / collection["id"]).exists()


def test_url_import_disabled_by_default(admin, collection):
    resp = admin.post(f"/collections/{collection['id']}/documents/url", json={"url": "https://eur-lex.europa.eu"})
    assert resp.status_code == 403


def test_url_import_rejects_internal_targets(admin, collection, monkeypatch):
    monkeypatch.setattr(get_settings(), "allow_url_ingestion", True)
    for url in ["http://127.0.0.1/admin", "http://localhost:8000/", "file:///etc/passwd", "http://10.0.0.5/"]:
        resp = admin.post(f"/collections/{collection['id']}/documents/url", json={"url": url})
        assert resp.status_code == 400, url


def test_reindex_legacy_document_without_file(admin, collection, db):
    legacy = Document(collection_id=uuid.UUID(collection["id"]), filename="ancien.pdf", status="failed")
    db.add(legacy)
    db.commit()
    assert admin.post(f"/documents/{legacy.id}/reindex").status_code == 409
