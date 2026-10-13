"""Format-specific text extraction.

Every loader returns ``PageText`` items. Headings found in structured formats (DOCX
styles, HTML ``<h1>``–``<h6>``) are emitted as Markdown ``#`` lines so the chunker
can attach a section label to each chunk regardless of the source format.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from src.config import get_settings
from src.ingestion.filetypes import SUPPORTED_TYPES, IngestionError, decode_text
from src.utils.logging_config import get_logger

log = get_logger()

# Below this many characters a PDF page is treated as having no text layer.
_MIN_PAGE_CHARS = 20


@dataclass
class PageText:
    number: int | None  # 1-based page number, None for unpaginated formats
    text: str


@dataclass
class LoadedDocument:
    pages: list[PageText]
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)
    ocr_pages: int = 0


@lru_cache
def ocr_available() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_pdf_page(path: Path, page_index: int) -> str:
    import pypdfium2 as pdfium
    import pytesseract

    pdf = pdfium.PdfDocument(str(path))
    try:
        page = pdf[page_index]
        image = page.render(scale=300 / 72).to_pil()
        return pytesseract.image_to_string(image, lang=get_settings().ocr_languages)
    finally:
        pdf.close()


def load_pdf(path: Path) -> LoadedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(""):
            raise IngestionError("PDF protégé par mot de passe : retirez la protection puis réimportez-le.")
    except PdfReadError as exc:
        raise IngestionError(f"PDF illisible : {exc}") from exc

    settings = get_settings()
    pages: list[PageText] = []
    empty: list[int] = []
    for index, page in enumerate(reader.pages):
        try:
            content = page.extract_text() or ""
        except Exception as exc:  # pypdf raises a zoo of errors on damaged pages
            log.warning("[ingestion] page %d illisible : %s", index + 1, exc)
            content = ""
        if len(content.strip()) < _MIN_PAGE_CHARS:
            empty.append(index)
        pages.append(PageText(index + 1, content))

    loaded = LoadedDocument(pages=pages, page_count=len(pages))
    if empty:
        if settings.ocr_enabled and ocr_available():
            for index in empty:
                try:
                    pages[index].text = _ocr_pdf_page(path, index)
                    loaded.ocr_pages += 1
                except Exception as exc:
                    log.warning("[ingestion] OCR échoué page %d : %s", index + 1, exc)
            loaded.warnings.append(f"{loaded.ocr_pages} page(s) numérisée(s) traitée(s) par OCR.")
        else:
            loaded.warnings.append(
                f"{len(empty)} page(s) sans couche texte ignorée(s) (OCR désactivé ou Tesseract absent)."
            )
    return loaded


def load_docx(path: Path) -> LoadedDocument:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise IngestionError(f"Document Word illisible : {exc}") from exc

    lines: list[str] = []
    for block in document.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = docx.text.paragraph.Paragraph(block, document)
            content = paragraph.text.strip()
            if not content:
                continue
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if style.startswith(("heading", "titre", "title")):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                lines.append(f"{'#' * min(int(level), 6)} {content}")
            else:
                lines.append(content)
        elif tag == "tbl":
            table = docx.table.Table(block, document)
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    return LoadedDocument(pages=[PageText(None, "\n".join(lines))], page_count=0)


def load_text(path: Path) -> LoadedDocument:
    return LoadedDocument(pages=[PageText(None, decode_text(path.read_bytes()))], page_count=0)


def load_html_bytes(raw: bytes) -> LoadedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(raw), "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "form"]):
        tag.decompose()
    for level in range(1, 7):
        for heading in soup.find_all(f"h{level}"):
            heading.replace_with(f"\n{'#' * level} {heading.get_text(' ', strip=True)}\n")
    content = soup.get_text("\n")
    lines = [line.strip() for line in content.splitlines()]
    return LoadedDocument(pages=[PageText(None, "\n".join(line for line in lines if line))], page_count=0)


def load_file(path: Path, mime_type: str) -> LoadedDocument:
    if mime_type == SUPPORTED_TYPES[".pdf"]:
        return load_pdf(path)
    if mime_type == SUPPORTED_TYPES[".docx"]:
        return load_docx(path)
    if mime_type == "text/html":
        return load_html_bytes(path.read_bytes())
    if mime_type in ("text/plain", "text/markdown"):
        return load_text(path)
    raise IngestionError(f"Type de fichier non supporté : {mime_type}")
